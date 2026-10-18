"""Fill the AI Builder Cup template with Canon's pitch deck.

  python3 deck/build_deck.py [--shots web/shots] [--mocks mocks/v2/shots] [--out deck/canon-deck.pptx]

Images are read from folders, so the lead can drop final screenshots in and rebuild.
Snapshot slide uses SNAPSHOTS (file names inside --shots); missing files are skipped.
Needs python-pptx and Pillow. Export to PDF with deck/export_pdf.sh (Keynote).
"""
import argparse
import copy
import json
import subprocess
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEAM_NAME = "Localhost"
LEADER = "Tejas Nafde"
PROBLEM = "Future of Work & Enterprise Productivity"
APP_URL = "https://canon-tn07.web.app"
RUN_URL = "https://canon-ezvux7gqxa-el.a.run.app"
DISCORD_URL = "https://discord.gg/HwSgjmcgF5"
REPO_URL = "https://github.com/tejasnafde/canon"
VIDEO_URL = "https://youtu.be/tPJenhhhjI8"
SNAPSHOTS = [("home-1080.png", "Home: four stories and live counts"), ("new-hire-1080.png", "I am new here: out-of-date page, who to ask"),
             ("catch-up-sales-1080.png", "Catch me up: sales, ranked to-do list"), ("agent-1080.png", "An AI agent gets stopped: Block"),
             ("ask-1080.png", "Ask why: cited answer and timeline"), ("org-1080.png", "People: org tree, who left, successors")]
LOGO = ROOT / "brand" / "final" / "canon-lockup.svg"

FONT = "Google Sans Flex"
# brand/final colours: ink, vermilion, paper
INK, MUTED, ACCENT = RGBColor(0x1C, 0x1A, 0x17), RGBColor(0x5F, 0x5A, 0x52), RGBColor(0xC8, 0x36, 0x1E)
FILL, LINE, PAPER, WHITE = RGBColor(0xF6, 0xF1, 0xE7), RGBColor(0xD8, 0xCF, 0xBE), RGBColor(0xFB, 0xF4, 0xF1), RGBColor(0xFF, 0xFF, 0xFF)
X0, Y0, W, H = Inches(0.35), Inches(1.3), Inches(9.3), Inches(4.05)  # content area under the title


def drop_slide(prs, index):
    ids = prs.slides._sldIdLst
    sid = ids[index]
    prs.part.drop_rel(sid.get(qn("r:id")))
    ids.remove(sid)


def style(run, size, bold=False, color=INK):
    run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FONT, Pt(size), bold, color


def set_title(slide, text):
    box = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip())
    tf = box.text_frame
    p = tf.paragraphs[0] if tf.paragraphs[0].text.strip() else next(p for p in tf.paragraphs if p.text.strip())
    for extra in [x for x in tf.paragraphs if x._p is not p._p]:
        extra._p.getparent().remove(extra._p)
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = text
    box.height = Inches(0.6)


def text(slide, x, y, w, h, items, size=12, color=INK, gap=4, align=None, anchor=MSO_ANCHOR.TOP):
    """items: str or (bold_head, rest). A leading '- ' makes a bullet-like indent."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, it in enumerate(items if isinstance(items, list) else [items]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        if align:
            p.alignment = align
        head, rest = it if isinstance(it, tuple) else (None, it)
        if head:
            style(p.add_run(), size, True, color)
            p.runs[-1].text = head + (" " if rest else "")
        if rest:
            r = p.add_run()
            style(r, size, False, color)
            r.text = rest
    return tb


def box(slide, x, y, w, h, label, size=10, fill=FILL, color=INK, bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = LINE
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        style(r, size if i == 0 else size - 2, bold if i == 0 else False, color if i == 0 else MUTED)
        r.text = line
    return s


def arrow(slide, a, b, side="h"):
    """Straight connector from shape a to shape b with an arrowhead at b."""
    if side == "h":
        x1, y1, x2, y2 = a.left + a.width, a.top + a.height // 2, b.left, b.top + b.height // 2
    else:
        x1, y1, x2, y2 = a.left + a.width // 2, a.top + a.height, b.left + b.width // 2, b.top
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb, c.line.width = MUTED, Pt(1.25)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def picture(slide, path, x, y, w, h, tmp):
    """Downscale to JPEG (keeps the PDF small), then fit inside the box, centred, with a thin border."""
    im = Image.open(path).convert("RGB")
    im.thumbnail((1400, 1400))
    out = Path(tmp) / (path.stem + ".jpg")
    im.save(out, quality=82)
    r = min(w / im.width, h / im.height)
    pw, ph = int(im.width * r), int(im.height * r)
    pic = slide.shapes.add_picture(str(out), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)
    pic.line.color.rgb, pic.line.width = LINE, Pt(0.75)
    return pic


def table(slide, x, y, w, rows, widths, size=10, row_h=Inches(0.28)):
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows)).table
    for j, cw in enumerate(widths):
        t.columns[j].width = Emu(int(w * cw))
    for i, row in enumerate(rows):
        t.rows[i].height = row_h
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = INK if i == 0 else (FILL if i % 2 else WHITE)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            r = p.add_run()
            style(r, size, i == 0, WHITE if i == 0 else INK)
            r.text = str(val)
    return t


def logo(slide, x, y, w, tmp):
    """The lockup SVG rasterised with rsvg-convert (python-pptx cannot place SVG)."""
    png = Path(tmp) / "lockup.png"
    if not png.exists():
        subprocess.run(["rsvg-convert", "-w", "900", str(LOGO), "-o", str(png)], check=True)
    return slide.shapes.add_picture(str(png), x, y, width=w)


def caption(slide, x, y, w, s):
    text(slide, x, y, w, Inches(0.3), s, size=9, color=MUTED, align=PP_ALIGN.CENTER)




def build(shots: Path, mocks: Path, out: Path):
    prs = Presentation(str(ROOT / "deck" / "template.pptx"))
    drop_slide(prs, 14)  # empty spare content slide
    drop_slide(prs, 1)   # "Note: Please download this template..." instruction slide
    s = list(prs.slides)
    ev = json.loads((ROOT / "backend" / "canon" / "eval_report.json").read_text())
    tmp = tempfile.mkdtemp()

    # 1. Team details
    tf = next(x for x in s[0].shapes if x.has_text_frame).text_frame
    for p, val in zip([p for p in tf.paragraphs if p.text.strip()][1:], [TEAM_NAME, LEADER, PROBLEM]):
        r = p.add_run()
        style(r, 18, True)
        r.text = val

    # 2. Brief
    set_title(s[1], "Brief about the idea")
    logo(s[1], Inches(8.25), Inches(0.8), Inches(1.4), tmp)
    text(s[1], X0, Y0, Inches(4.5), H, [
        ("Canon is the shared memory of a company, for its people and its AI agents.", ""),
        "Teams decide things in Slack threads, meetings and pull requests, then forget. Wiki pages go out of date. New hires and coding agents act on old rules.",
        "Canon keeps one record of decisions: rationale, owner, date, authority, and a quote checked word for word against its source.",
        "It checks every new message, document and agent action against that record, and warns when one is about to break a decision.",
        "Coding agents connect over MCP and call check_intent before they act. A Friday deploy job is blocked, with the incident review as evidence.",
        "Live on a public Discord server: post a message, and Canon replies within a minute.",
    ], size=11, gap=6)
    if (shots / "agent-1080.png").exists():
        pic = picture(s[1], shots / "agent-1080.png", Inches(5.0), Y0, Inches(4.65), Inches(3.6), tmp)
        caption(s[1], Inches(5.0), pic.top + pic.height, Inches(4.65), "canon.tn07.dev: an AI agent gets stopped")

    # 3. Opportunities
    set_title(s[2], "Opportunities")
    cols = [
        ("How is it different?", [
            "Meeting-memory tools summarise. Decision logs store. Spec checkers review code.",
            "Canon does all three on one record, and knows who may change a rule: official, proposed or claim, from the org tree.",
            "Its users include AI agents, not only people."]),
        ("How it solves the problem", [
            "Extracts decisions with verified quotes from chats, meetings, wiki, PRs, PDFs, screenshots and voice recordings.",
            "Reconciles each one: same, update or contradiction, ordered by effective date.",
            "Flags out-of-date documents and broken premises, and sends unverified claims to the right lead.",
            "Answers why with citations, or says I do not know."]),
        ("USP: agent-native", [
            "An MCP server any agent can call: check_intent returns allow, warn or block with the deciding decision.",
            "Every agent read, write and check is recorded for people to see.",
            "Also an ADK agent over A2A, and a Discord bot."]),
    ]
    cw = Inches(3.0)
    for i, (head, body) in enumerate(cols):
        x = X0 + i * (cw + Inches(0.15))
        box(s[2], x, Y0, cw, Inches(0.45), head, size=13, fill=INK, color=WHITE, bold=True, shape=MSO_SHAPE.RECTANGLE)
        text(s[2], x, Y0 + Inches(0.55), cw, Inches(3.4), ["- " + b for b in body], size=11, gap=6)

    # 4. Features
    set_title(s[3], "List of features offered by the solution")
    left = [("Ingest anything:", "Slack, meetings, wiki, PRs, commits, email, paste, PDFs, screenshots, audio and video (Gemini reads them; 10 min audio, 2 min video, 30 min a day)."),
            ("Decision extraction:", "statement, rationale, assumptions, owner, date, quotes verified word for word."),
            ("Reconciliation:", "same, update (supersedes) or contradiction, ordered by effective date. Older copies marked outdated."),
            ("Authority:", "official, proposed or claim from the org tree. Unverified claims go to the right lead. Jokes are ignored."),
            ("Drift alerts:", "out-of-date documents and broken premises, both quotes side by side."),
            ("Ask why:", "cited answer and timeline; refuses when the record is silent.")]
    right = [("I am new here:", "five things you must know, a page not to trust, people to ask."),
             ("Catch me up:", "14 roles, a summary and a ranked to-do list with reasons."),
             ("Agent guard over MCP:", "check_intent gives allow, warn or block. Six tools for Claude Code and Codex. ADK agent over A2A."),
             ("Discord, live:", "synced every minute (threads, files, PDFs, screenshots, voice notes); replies under the message; /canon ask, check, catchup, remember."),
             ("People and org:", "who owns what, who left, who took over."),
             ("Measured:", "gold-set benchmark and a memory-health replay, served in the app."),
             ("Cost controls:", "daily Gemini cap, Discord caps, scale to zero.")]
    text(s[3], X0, Y0, Inches(4.55), H, left, size=10.5, gap=6)
    text(s[3], X0 + Inches(4.75), Y0, Inches(4.55), H, right, size=10.5, gap=6)

    # 5. Process flow
    set_title(s[4], "Process flow diagram")
    bw, bh, gx = Inches(1.38), Inches(0.75), Inches(0.2)
    text(s[4], X0, Y0 - Inches(0.05), Inches(6), Inches(0.3), ("A. A new source arrives (web paste, Discord, PR, agent, PDF)", ""), size=11)
    row1 = ["Source\ntext, PDF or image", "Chunk + embed\ngemini-embedding-001", "Extract decisions\nGemini, structured", "Verify quotes\nexact substring",
            "Authority + reconcile\nsame / update / contra", "Drift scan\nold docs, premises"]
    y1 = Y0 + Inches(0.32)
    a = [box(s[4], X0 + i * (bw + gx), y1, bw, bh, t) for i, t in enumerate(row1)]
    for i in range(len(a) - 1):
        arrow(s[4], a[i], a[i + 1])
    rec = box(s[4], X0 + Inches(1.9), y1 + Inches(1.15), Inches(4.3), Inches(0.6),
              "Firestore record: decisions, alerts, activity, people\nvector search, workspace prefilter", size=11, fill=PAPER, bold=True)
    arrow(s[4], a[4], rec, "v")
    arrow(s[4], a[5], rec, "v")  # the drift scan writes stale-doc and premise alerts into the record
    text(s[4], X0, Y0 + Inches(3.5), Inches(6), Inches(0.3), ("B. A person or an agent asks or acts", ""), size=11)
    row2 = ["Question, briefing\nor intended action", "Embed query\nnearest decisions", "Below threshold?\nrefuse / allow, no LLM",
            "Gemini judges\nanswer or verdict", "Cite verified quotes\nallow / warn / block", "Recorded\nweb, MCP, A2A, Discord"]
    y2 = Y0 + Inches(2.67)
    b = [box(s[4], X0 + i * (bw + gx), y2, bw, bh, t) for i, t in enumerate(row2)]
    for i in range(len(b) - 1):
        arrow(s[4], b[i], b[i + 1])
    c = arrow(s[4], rec, b[1], "v")
    c.begin_x = c.end_x = b[1].left + b[1].width // 2  # straight down into "Embed query"

    # 6. Wireframes (the v2 HTML mocks the web app was built from)
    set_title(s[5], "Wireframes / mock diagrams")
    text(s[5], X0, Y0 + Inches(3.95), W, Inches(0.3), "Early wireframes: the approved mocks the web app was built from. The live app (current numbers) is on the Snapshots slide.", size=10, color=MUTED)
    names = [("home.png", "Home: stories and live counts"), ("new-hire.png", "I am new here"), ("agent.png", "For agents: the block")]
    spots = [(X0, Y0, Inches(5.6), Inches(3.5)), (X0 + Inches(5.8), Y0, Inches(3.5), Inches(1.65)),
             (X0 + Inches(5.8), Y0 + Inches(1.95), Inches(3.5), Inches(1.65))]
    for (n, cap), (x, y, w, h) in zip(names, spots):
        if (mocks / n).exists():
            pic = picture(s[5], mocks / n, x, y, w, h, tmp)
            caption(s[5], x, pic.top + pic.height, w, cap)

    # 7. Architecture
    set_title(s[6], "Architecture diagram")
    cx = [box(s[6], X0, Y0 + i * Inches(0.95), Inches(1.9), Inches(0.75), t, size=10) for i, t in enumerate([
        "Web app (React)\nFirebase Hosting", "Claude Code, Codex\nMCP clients", "Other agents\nA2A", "Discord server\nCloud Scheduler, 1 min"])]
    run = s[6].shapes.add_shape(MSO_SHAPE.RECTANGLE, X0 + Inches(2.55), Y0 - Inches(0.05), Inches(3.4), H)
    run.fill.solid(); run.fill.fore_color.rgb = PAPER; run.line.color.rgb = ACCENT
    run.shadow.inherit = False
    text(s[6], run.left, run.top + Inches(0.02), run.width, Inches(0.3), ("Cloud Run: canon, asia-south1, min 0", ""), size=10,
         color=ACCENT, align=PP_ALIGN.CENTER)
    ep = [box(s[6], run.left + Inches(0.2), Y0 + Inches(0.35) + i * Inches(0.62), Inches(1.3), Inches(0.5), t, size=10, fill=WHITE)
          for i, t in enumerate(["FastAPI /api", "MCP /mcp", "ADK agent /a2a", "Discord sync"])]
    core = box(s[6], run.left + Inches(1.85), Y0 + Inches(0.35), Inches(1.35), Inches(2.36),
               "Core\ningest, authority\nreconcile, ask\nbriefing, check", size=11, fill=WHITE, bold=True)
    for c, e in zip(cx, ep):
        arrow(s[6], c, e)
    for e in ep:
        arrow(s[6], e, core)
    box(s[6], run.left + Inches(0.2), Y0 + Inches(2.95), Inches(3.0), Inches(0.75),
        "Guardrails\nverified quotes, refusal threshold, source fencing,\ninput limits, timeouts, daily Gemini cap", size=10, fill=WHITE)
    gcp = [box(s[6], X0 + Inches(6.55), Y0 + i * Inches(1.3), Inches(2.75), Inches(1.0), t, size=11) for i, t in enumerate([
        "Vertex AI Gemini\ngemini-3.5-flash-lite\nstructured output, PDFs, images", "Embeddings\ngemini-embedding-001, 768 dims",
        "Firestore (database canon)\nvector search (cosine)\nworkspace prefilter"])]
    for g in gcp:
        arrow(s[6], core, g)

    # 8. Technologies
    set_title(s[7], "Technologies used in the solution")
    table(s[7], X0, Y0, W, [
        ["Technology", "What Canon uses it for"],
        ["Gemini 3.5 Flash-Lite on Vertex AI", "Extraction, reconciliation, drift checks, answers, briefings, verdicts; structured output; reads PDFs and images"],
        ["gemini-embedding-001 (768 dims)", "Embeddings for source chunks, decisions and queries"],
        ["Cloud Firestore, named database canon", "Sources, chunks, decisions, alerts, activity, people; vector search with a workspace prefilter"],
        ["Cloud Run", "One scale-to-zero service: API, MCP server, A2A agent, Discord sync and slash commands"],
        ["Firebase Hosting", "Static web app; rewrites /api, /mcp, /a2a, /hooks, /discord to Cloud Run on one origin"],
        ["Cloud Scheduler", "Calls the Discord sync every minute, so no always-on gateway is needed"],
        ["Secret Manager, Cloud Monitoring", "One secret for every Canon credential; alert policies relayed to Discord"],
        ["ADK + A2A, MCP (FastMCP)", "Canon as an LlmAgent with four tools; six MCP tools over streamable HTTP"],
        ["Python 3.12, FastAPI, uv, pytest", "Backend; 85 tests replay recorded Gemini fixtures, no network"],
        ["React, TypeScript, Vite, Tailwind", "Editorial web app: stories, catch-up, agents, paste, ask, decisions, people, benchmark"],
    ], [0.33, 0.67], size=9.5, row_h=Inches(0.35))

    # 9. Cost
    set_title(s[8], "Estimated implementation cost")
    text(s[8], X0, Y0, Inches(4.55), H, [
        ("Idle cost is near zero.", "Cloud Run scales to zero (min instances 0). No per-hour resource."),
        ("Firestore:", "the demo data fits inside the free tier."),
        ("Gemini is paid per call.", "An answer is one embedding and one Flash-Lite call. A check is one call, or none when no decision is relevant. An ingest is one extraction plus a few reconcile calls."),
        ("Refusals and irrelevant checks never reach Gemini.", "The similarity threshold stops them first."),
        ("Tracked, not guessed:", "Gemini calls, Cloud Run, Artifact Registry and Hosting carry the label app=canon, so the billing export splits Canon's spend."),
    ], size=11, gap=7)
    text(s[8], X0 + Inches(4.8), Y0, Inches(4.5), H, [
        ("What we measured:", ""),
        "- Reading the whole 380-source archive (a full seed) costs about INR 80 to 100 in Gemini calls.",
        "- Live checks and answers: one call each, about 1 to 2 seconds.",
        ("Hard caps:", ""),
        "- Daily Gemini cap per instance (1500 generate calls by default).",
        "- Discord: 200 messages a day, 20 per user; slash commands stop at the cap.",
        "- 120 POST requests per minute per IP.",
        ("Not measured yet:", "tokens per call at scale, so we do not quote a per-user price."),
    ], size=11, gap=5)

    # 10. Snapshots: 3 x 2 grid of the final 1080p screenshots
    set_title(s[9], "Snapshots of the prototype")
    found = [(n, c) for n, c in SNAPSHOTS if (shots / n).exists()]
    sw, sh = Inches(2.98), Inches(1.68)
    for i, (n, cap) in enumerate(found[:6]):
        x, y = X0 + (i % 3) * (sw + Inches(0.18)), Y0 + (i // 3) * (sh + Inches(0.36))
        picture(s[9], shots / n, x, y, sw, sh, tmp)
        caption(s[9], x, y + sh, sw, cap)

    # 10b. Live in Discord: real screenshots of the public demo server, placed right after Snapshots
    disc = sorted((ROOT / "deck" / "discord").glob("*.png"))
    if disc:
        d = prs.slides.add_slide(s[9].slide_layout)
        ids = prs.slides._sldIdLst
        el = ids[-1]; ids.remove(el); ids.insert(10, el)  # after Snapshots (index 9)
        # the bare layout has no title box: clone the Snapshots title (same position and style)
        src = next(x for x in s[9].shapes if x.has_text_frame and x.text_frame.text.strip())
        d.shapes._spTree.append(copy.deepcopy(src._element))
        set_title(d, "Live in Discord: Canon learns from a real server")
        caps = {"1-general": "Threads in #general, posted as the people who wrote them",
                "2-engineering": "#engineering: the stale Friday release window, as a screenshot",
                "3-sales": "#sales: the old USD 99 pricing page and the sales FAQ file",
                "0-reply": "A real post at 21:14; Canon replied in about a minute. Join: discord.gg/HwSgjmcgF5"}
        dw, dh = Inches(4.55), Inches(1.75)
        for i, f in enumerate(disc[:4]):
            x = X0 + (i % 2) * (dw + Inches(0.2)); y = Y0 + (i // 2) * (dh + Inches(0.45))
            picture(d, f, x, y, dw, dh, tmp)
            caption(d, x, y + dh, dw, caps.get(f.stem, ""))

    # 11. Benchmark
    set_title(s[10], "Prototype performance report / benchmarking")
    pct = lambda a, b: f"{round(100 * a / b)}%" if b else ""
    note = {n.split(" found ")[0]: n.split(" found ")[1] for n in ev["notes"] if " found " in n and n.count("found") == 1}
    stale = note.get("stale docs", "").split("/")
    table(s[10], X0, Y0, Inches(5.2), [
        ["Task (labelled gold set)", "Planted", "Found", "Score"],
        ["Decisions extracted", ev["decisions_planted"], ev["decisions_found"], pct(ev["decisions_found"], ev["decisions_planted"])],
        ["Contradictions and updates", ev["contradictions_planted"], ev["contradictions_found"], pct(ev["contradictions_found"], ev["contradictions_planted"])],
        ["False positive alerts", "", ev["false_positives"], ""],
        ["Out-of-date documents", stale[1], stale[0], pct(int(stale[0]), int(stale[1]))],
        ["check_intent verdicts", ev["check_cases"], ev["check_correct"], pct(ev["check_correct"], ev["check_cases"])],
        ["Unanswerable, refused", ev["unanswerable"], ev["refused_correctly"], pct(ev["refused_correctly"], ev["unanswerable"])],
        ["Latency p50 / p95 (ask, check)", "", f"{ev['p50_ms']} / {ev['p95_ms']} ms", ""],
    ], [0.46, 0.13, 0.29, 0.12], size=10, row_h=Inches(0.34))
    text(s[10], X0, Y0 + Inches(2.85), Inches(5.2), Inches(1.2), [
        f"Run {ev['run_at'][:16].replace('T', ' ')} UTC, model {ev['model']}, dataset acme4: 380 synthetic sources (about 100 pure noise), 14 roles. "
        "check_intent re-measured 2026-10-18.",
        "Reproduce: backend/scripts/seed.py, then eval.py. The app serves the report at /api/eval."],
        size=9, color=MUTED, gap=4)
    rx = X0 + Inches(5.45)
    text(s[10], rx, Y0, Inches(3.85), Inches(0.3), ("Memory health: replay in time order", ""), size=11)
    table(s[10], rx, Y0 + Inches(0.35), Inches(3.85), [
        ["Check", "Result"],
        ["Decisions with no traceable source", f"{ev['health_untraceable']} of {ev['health_decisions_stored']}"],
        ["Duplicate decisions", f"{ev['health_duplicates']}"],
        ["Quarter checkpoints correct", f"{ev['health_checkpoints_correct']} of {ev['health_checkpoints']}"],
        ["Supersession chains correct", f"{ev['health_chains_correct']} of {ev['health_chains']}"],
        ["Authority cases correct", f"{ev['authority_correct']} of {ev['authority_cases']}"],
    ], [0.7, 0.3], size=9.5, row_h=Inches(0.3))
    text(s[10], rx, Y0 + Inches(2.3), Inches(3.85), Inches(1.7), [("In the open:", "")] + ["- " + n for n in [
        "premise changes found 2/5, our weakest part",
        "updates 19/20, contradictions 16/19",
        "answerable questions 21/21 cited an expected decision",
        "579 stored quotes, all verified word for word",
        "check misses lean to caution (warn)"]], size=9, gap=2)

    # 12. Future
    set_title(s[11], "Additional details / future development")
    text(s[11], X0, Y0, Inches(4.55), H, [
        ("Where people already talk:", "Slack and Microsoft Teams next to the live Discord connector, answering in place."),
        ("Permissions:", "sign-in, per-team visibility, and decisions that follow the access rules of their source."),
        ("Code review at scale:", "check PR diffs against the record and post review comments for every repo."),
    ], size=12, gap=9)
    text(s[11], X0 + Inches(4.8), Y0, Inches(4.5), H, [
        ("Beyond software:", "manufacturing SOP drift. A line change or a supplier note that contradicts the approved procedure raises an alert before a shift runs it."),
        ("Better recall:", "premise watch finds 2 of 5 today; we will tune it against the gold set."),
        ("Real data:", "this repository's own README, PRs and commits next to the synthetic corpus."),
    ], size=12, gap=9)

    # 13. Links
    set_title(s[12], "Links")
    links = text(s[12], X0, Y0, W, H, [
        ("GitHub public repository:", REPO_URL),
        ("Demo video:", VIDEO_URL),
        ("Final product link:", APP_URL),
        ("Same app, custom domain:", "https://canon.tn07.dev"),
        ("Cloud Run:", RUN_URL + "/health"),
        ("Live Discord demo:", DISCORD_URL),
        ("Connect your agent:", "claude mcp add --transport http canon https://canon.tn07.dev/mcp"),
        ("", "No login needed. All company data is synthetic and labelled as such."),
    ], size=14, gap=10)
    for p in links.text_frame.paragraphs:  # clickable URLs, same ink colour
        for r in p.runs:
            if r.text.startswith("https://"):
                r.hyperlink.address = r.text.split()[0]
                r.font.color.rgb = INK
    logo(s[12], Inches(7.6), Inches(4.7), Inches(2.0), tmp)

    prs.save(str(out))
    print("wrote", out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--shots", default=str(ROOT / "web" / "shots"))
    ap.add_argument("--mocks", default=str(ROOT / "mocks" / "v2" / "shots"))
    ap.add_argument("--out", default=str(ROOT / "deck" / "canon-deck.pptx"))
    a = ap.parse_args()
    build(Path(a.shots), Path(a.mocks), Path(a.out))
